from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

#
from pdf_toolkits.pdf_text_processing import PDF_text


def get_pdf_para_text(pdf_fle):
    '''to transfer defined pdf object into module and para to estimate'''
    pdf_object = PDF_text(pdf_fle)
    pdf_sub_modules = pdf_object.sub_module
    para_list = pdf_object.para_list
    return para_list

def get_section_contents(pdf_fle,subtitle_name):
    pdf_object = PDF_text(pdf_fle)
    sub_modules = pdf_object.sub_module
    for sub_module in sub_modules:
        subtitle = sub_module.subtitle
        if subtitle_name == subtitle:
            contents = sub_module.module_contents
            contents.insert(0,subtitle)
    return contents

from docx import Document
def read_sections_from_docx(file_path):
    doc = Document(file_path)

    sections = []  # List of sections
    current_section = []  # Current section (list of paragraphs)
    blank_line_count = 0

    for para in doc.paragraphs:
        text = para.text.strip()

        if not text:
            blank_line_count += 1
            continue
        else:
            if blank_line_count >= 2:
                # Two or more blank lines = new section
                if current_section:
                    sections.append(current_section)
                    current_section = []
            blank_line_count = 0  # reset counter on actual text
            current_section.append(text)

    # Append last section if exists
    if current_section:
        sections.append(current_section)

    return sections

from docx import Document
def read_paragraphs_from_docx(file_path):
    doc = Document(file_path)
    paragraphs = [para.text.strip() for para in doc.paragraphs if para.text.strip()]
    return paragraphs



from sklearn.metrics.pairwise import cosine_similarity
from sentence_transformers import SentenceTransformer
from bert_score import score
import numpy as np

def aligned_bert_score(extracted_texts, ground_truth_texts, lang="en", model_name="all-MiniLM-L6-v2", verbose=False):
    """
    Aligns extracted texts to the most similar ground truth texts using cosine similarity,
    then computes BERTScore for the aligned pairs.

    Parameters:
    - extracted_texts: list of str — the extracted paragraphs/texts
    - ground_truth_texts: list of str — the reference paragraphs/texts
    - lang: language code for BERTScore (default 'en')
    - model_name: sentence transformer model to use for alignment
    - verbose: whether to print progress info

    Returns:
    - precision, recall, f1: mean BERTScore values
    """

    if verbose:
        print(f"Loading embedding model: {model_name}")

    model = SentenceTransformer(model_name)

    if verbose:
        print("Encoding extracted and ground truth texts...")

    ext_emb = model.encode(extracted_texts)
    gt_emb = model.encode(ground_truth_texts)

    sim_matrix = cosine_similarity(ext_emb, gt_emb)

    # Greedy alignment: for each extracted text, find most similar ground truth
    best_matches = np.argmax(sim_matrix, axis=1)
    aligned_gt = [ground_truth_texts[i] for i in best_matches]

    if verbose:
        print("Computing BERTScore...")

    P, R, F1 = score(extracted_texts, aligned_gt, lang=lang, verbose=verbose)

    return P.mean().item(), R.mean().item(), F1.mean().item()


### text use troch should be not in InfoExtractor environment
gt_file =r"D:\All_self_files\info_extract_examples\621 pdf test\digital\nature_synthesis\ns_5\ground_truth.docx"
gt_paras =  read_paragraphs_from_docx(gt_file)

pdf_file = r"D:\All_self_files\info_extract_examples\621 pdf test\digital\nature_synthesis\ns_5\s44160-025-00793-9.pdf"
extracted_text=get_pdf_para_text(pdf_file)
P,R,F1 = aligned_bert_score(gt_paras,extracted_text)
print(P,R,F1)
